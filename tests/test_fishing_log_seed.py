"""엑셀 초기 이관 파서 (spec §3). 실제 엑셀은 커밋하지 않는다 - 값은 전부
지어낸 것이고, 필요한 엑셀은 테스트 안에서 openpyxl로 만든다."""
from services.fishing_log.excel_seed import (
    extract_ship_name, normalize_category, normalize_shop, parse_catches,
    ship_match_key, to_int,
)


def test_to_int_accepts_numbers_and_digit_strings():
    assert to_int(12000) == 12000
    assert to_int(12000.0) == 12000
    assert to_int(1234.6) == 1235
    assert to_int('12,000') == 12000


def test_to_int_returns_none_for_text_and_blanks():
    assert to_int('11만원') is None
    assert to_int('무료') is None
    assert to_int(None) is None
    assert to_int('') is None
    assert to_int(True) is None


def test_normalize_shop_merges_known_spellings():
    assert normalize_shop('테무') == 'Temu'
    assert normalize_shop('temu') == 'Temu'
    assert normalize_shop('aliexpress') == 'AliExpress'
    assert normalize_shop('Aliexpress') == 'AliExpress'
    assert normalize_shop('에프마켓 인천') == '에프마켓인천'
    assert normalize_shop('  가나낚시 ') == '가나낚시'
    assert normalize_shop(None) is None


def test_normalize_category_maps_synonyms_and_blank():
    assert normalize_category('줄') == '라인'
    assert normalize_category('낚시대') == '로드'
    assert normalize_category('에기') == '에기'
    assert normalize_category(None) == '기타'


def test_ship_match_key_ignores_spaces_case_and_trailing_ho():
    assert ship_match_key('가나다호') == ship_match_key(' 가나 다 ')
    assert ship_match_key('ABC호') == ship_match_key('abc')


def test_extract_ship_name_prefers_word_ending_in_ho():
    assert extract_ship_name('가나다호 - (농어, 우럭)') == ('가나다호', '(농어, 우럭)')
    assert extract_ship_name('라마낚시 바사호 문어') == ('바사호', '라마낚시 문어')
    assert extract_ship_name('테스트도 원강호(다른호)') == ('원강호', '테스트도 (다른호)')
    assert extract_ship_name('하늘폭스') == ('하늘폭스', None)
    assert extract_ship_name('가나다 호') == ('가나다 호', None)


def test_parse_catches_simple_forms():
    assert parse_catches('쭈 10, 갑 2') == (
        [{'who': '나', 'species': '쭈꾸미', 'count': 10},
         {'who': '나', 'species': '갑오징어', 'count': 2}], False)
    assert parse_catches('문어3마리') == ([{'who': '나', 'species': '문어', 'count': 3}], False)


def test_parse_catches_tracks_people():
    catches, leftover = parse_catches('마눌 문5, 나 문7')
    assert catches == [{'who': '마눌', 'species': '문어', 'count': 5},
                       {'who': '나', 'species': '문어', 'count': 7}]
    assert leftover is False

    catches, _ = parse_catches('마눌 쭈4,갑1')
    assert [c['who'] for c in catches] == ['마눌', '마눌']

    catches, _ = parse_catches('가: 쭈3, 나: 쭈5')
    assert [(c['who'], c['count']) for c in catches] == [('가', 3), ('나', 5)]


def test_parse_catches_ignores_percent_and_weight_and_flags_leftover_digits():
    catches, leftover = parse_catches('쭈 10, 입질 60~70%')
    assert catches == [{'who': '나', 'species': '쭈꾸미', 'count': 10}]
    assert leftover is True

    catches, leftover = parse_catches('쭈70%')
    assert catches == [] and leftover is True


def test_parse_catches_no_digits_is_not_a_warning():
    assert parse_catches('꽝') == ([], False)
    assert parse_catches('배 깨끗함') == ([], False)
    assert parse_catches(None) == ([], False)


def test_parse_catches_does_not_treat_na_inside_words_as_person():
    catches, _ = parse_catches('나름 괜찮음 쭈 3')
    assert catches == [{'who': '나', 'species': '쭈꾸미', 'count': 3}]


from datetime import date, datetime

from openpyxl import Workbook

from services.fishing_log.excel_seed import YEAR_HEADER, parse_workbook

TODAY = date(2026, 6, 15)


def _book(year_rows=None, ship_rows=None, extra_sheets=None, header=None):
    """가짜 엑셀. year_rows: 2026 시트의 데이터 줄(10칸 리스트)."""
    wb = Workbook()
    ws = wb.active
    ws.title = '2026'
    ws.append(header or YEAR_HEADER)
    for row in year_rows or []:
        ws.append(row)
    if ship_rows is not None:
        ships = wb.create_sheet('낚시배')
        for row in ship_rows:
            ships.append(row)
    for title, rows in (extra_sheets or {}).items():
        sheet = wb.create_sheet(title)
        for row in rows:
            sheet.append(row)
    return wb


def _row(d=None, who=None, port=None, region=None, content=None, species=None,
         kind=None, note=None, price=None, catch=None):
    return [d, who, port, region, content, species, kind, note, price, catch]


def test_trip_row_becomes_done_trip_with_ship():
    wb = _book([_row(datetime(2026, 5, 2), '솔로', '가항', '가시', '가나다호 - (쭈갑)',
                     '쭈꾸미, 갑오징어', '선비', '배 깨끗', 90000, '쭈 10, 갑 2')])
    result = parse_workbook(wb, TODAY)

    assert len(result.trips) == 1
    trip = result.trips[0]
    assert trip.trip_date == date(2026, 5, 2)
    assert trip.status == 'done'
    assert trip.ship_name == '가나다호'
    assert trip.cost == 90000
    assert trip.companions == '솔로'
    assert trip.species == ['쭈꾸미', '갑오징어']
    assert trip.catches == [{'who': '나', 'species': '쭈꾸미', 'count': 10},
                            {'who': '나', 'species': '갑오징어', 'count': 2}]
    assert trip.catch_raw == '쭈 10, 갑 2'
    assert trip.memo == '(쭈갑)\n배 깨끗'
    assert [(s.name, s.region, s.port) for s in result.ships] == [('가나다호', '가시', '가항')]


def test_future_trip_is_planned_and_kind_bae_is_trip():
    wb = _book([_row(datetime(2026, 7, 1), '마눌', '나항', '나시', '라마호', '문어', '배', None, 200000)])
    assert parse_workbook(wb, TODAY).trips[0].status == 'planned'


def test_cancelled_rows_including_typo():
    wb = _book([
        _row(datetime(2026, 4, 5), '솔로', '가항', '가시', '바사호', '우럭', None, '날씨로 취소'),
        _row(datetime(2026, 4, 12), '솔로', '가항', '가시', '아자호', None, None, '최소됨'),
    ])
    trips = parse_workbook(wb, TODAY).trips
    assert [(t.ship_name, t.status) for t in trips] == [('바사호', 'cancelled'), ('아자호', 'cancelled')]


def test_blank_date_carries_previous_date_for_purchases():
    wb = _book([
        _row(datetime(2026, 3, 1), content='테스트 에기 3개', kind='에기', note='테무', price=12000),
        _row(None, content='테스트 봉돌', kind='봉돌', note='테무', price='3,000'),
    ])
    purchases = parse_workbook(wb, TODAY).purchases
    assert [(p.purchase_date, p.item, p.category, p.shop, p.price) for p in purchases] == [
        (date(2026, 3, 1), '테스트 에기 3개', '에기', 'Temu', 12000),
        (date(2026, 3, 1), '테스트 봉돌', '봉돌', 'Temu', 3000),
    ]


def test_total_rows_and_price_only_rows_are_skipped():
    wb = _book([
        _row(datetime(2026, 3, 1), content='테스트 줄', kind='줄', note='가나낚시', price=5000),
        _row('합계', price=5000),
        _row(None, price=5000),
    ])
    result = parse_workbook(wb, TODAY)
    assert [p.category for p in result.purchases] == ['라인']
    assert len(result.skipped) == 2


def test_first_row_without_date_is_warned_and_skipped():
    wb = _book([_row(None, content='테스트 에기', kind='에기', price=1000)])
    result = parse_workbook(wb, TODAY)
    assert result.purchases == []
    assert any('이어받을 일자' in w.reason for w in result.warnings)


def test_text_date_is_warned_and_skipped():
    wb = _book([
        _row(datetime(2026, 3, 1), content='첫 품목', kind='에기', price=1000),
        _row('2026.03.05', content='글자 날짜 품목', kind='에기', price=1000),
        _row(None, content='이어받는 품목', kind='에기', price=1000),
    ])
    result = parse_workbook(wb, TODAY)
    assert [(p.item, p.purchase_date) for p in result.purchases] == [
        ('첫 품목', date(2026, 3, 1)), ('이어받는 품목', date(2026, 3, 1))]
    assert any('날짜로 읽지 못함' in w.reason for w in result.warnings)


def test_row_without_kind_or_cancel_becomes_etc_purchase_with_warning():
    wb = _book([_row(datetime(2026, 3, 1), content='분류 없는 물건', note='가나낚시', price=2000)])
    result = parse_workbook(wb, TODAY)
    assert [(p.item, p.category) for p in result.purchases] == [('분류 없는 물건', '기타')]
    assert any('기타' in w.reason for w in result.warnings)


def test_trip_price_missing_and_unrecognized_catch_warns():
    wb = _book([_row(datetime(2026, 5, 2), '솔로', '가항', '가시', '가나다호', '문어', '선비',
                     None, None, '나(12), 마눌(3)')])
    result = parse_workbook(wb, TODAY)
    trip = result.trips[0]
    assert trip.cost is None
    assert trip.catches == []
    assert trip.catch_raw == '나(12), 마눌(3)'
    assert any('조과' in w.reason for w in result.warnings)


def test_ship_map_overrides_extraction():
    wb = _book([_row(datetime(2026, 5, 2), '솔로', '가항', '가시', '25시 수평선호 (새벽)', '쭈꾸미', '선비')])
    result = parse_workbook(wb, TODAY, ship_map={'25시 수평선호 (새벽)': '25시 수평선호'})
    assert result.trips[0].ship_name == '25시 수평선호'
    # 보정표로 선사명을 정해도 원문의 나머지 정보(출항 시각 등)는 잃지 않는다
    assert result.trips[0].memo == '25시 수평선호 (새벽)'


def test_header_mismatch_skips_sheet():
    wb = _book([_row(datetime(2026, 5, 2), content='x', kind='에기', price=1)],
               header=['날짜', '내용'])
    result = parse_workbook(wb, TODAY)
    assert result.purchases == [] and result.trips == []
    assert any('헤더' in w.reason for w in result.warnings)


def test_ships_sheet_carries_region_port_and_reads_travel_time():
    wb = _book(ship_rows=[
        ['가시', None, '가항', '가나다호', '문어 전문'],
        [None, None, None, '라마호'],
        [None, None, '나항', '바사호', '50분', '추천받음'],
        ['나시', None, None, '아자호', '1시간 7분'],
        ['무시할 줄', None, None, None],
    ])
    result = parse_workbook(wb, TODAY)
    ships = {s.name: s for s in result.ships}
    assert (ships['가나다호'].region, ships['가나다호'].port, ships['가나다호'].memo) == ('가시', '가항', '문어 전문')
    assert (ships['라마호'].region, ships['라마호'].port) == ('가시', '가항')
    assert (ships['바사호'].port, ships['바사호'].travel_time, ships['바사호'].memo) == ('나항', '50분', '추천받음')
    assert (ships['아자호'].region, ships['아자호'].port, ships['아자호'].travel_time) == ('나시', None, '1시간 7분')
    assert any('선사명' in w.reason for w in result.warnings)


def test_duplicate_ship_names_are_merged():
    wb = _book(ship_rows=[
        ['가시', None, '가항', '빅보스호', '메모1'],
        ['나시', None, '나항', '빅보스 호', '메모2'],
    ])
    ships = parse_workbook(wb, TODAY).ships
    assert len(ships) == 1
    assert ships[0].name == '빅보스호'
    assert ships[0].region == '가시'
    assert ships[0].memo == '메모1\n메모2'


def test_trip_ship_matches_sheet_ship_by_normalized_name():
    wb = _book(
        [_row(datetime(2026, 5, 2), '솔로', '딴항', '딴시', '가나다 호', '문어', '선비')],
        ship_rows=[['가시', None, '가항', '가나다호', '메모']],
    )
    result = parse_workbook(wb, TODAY)
    assert [s.name for s in result.ships] == ['가나다호']
    assert result.ships[0].port == '가항'
    assert result.trips[0].ship_name == '가나다호'


def test_gear_sheets_with_content_become_gear_items_and_map_is_ignored():
    wb = _book(extra_sheets={
        '테스트 릴': [['기어비 5.6', '우핸들'], ['합사 1호']],
        '빈 장비': [],
        '지도': [['무시']],
    })
    items = parse_workbook(wb, TODAY).gear_items
    assert [(g.name, g.memo) for g in items] == [('테스트 릴', '기어비 5.6 우핸들\n합사 1호')]


def test_row_with_companions_but_no_kind_is_a_trip():
    """실제 엑셀에서 품목 칸만 비운 출조 줄이 있었다 - 장비 구매 줄은 '함께'를
    채우지 않으므로, 품목이 비어도 함께가 있으면 출조로 본다."""
    wb = _book([_row(datetime(2026, 5, 2), '동출', None, None, '오천 가나다호 쭈갑', None, None, None, 0, '동출')])
    result = parse_workbook(wb, TODAY)
    assert result.purchases == []
    assert [(t.ship_name, t.status, t.companions) for t in result.trips] == [('가나다호', 'done', '동출')]
